from __future__ import annotations

from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.exceptions import ExportError
from app.core.logging import get_logger, log_duration
from app.schemas.lead import LeadResponse

logger = get_logger(__name__)

HEADERS = [
    "First Name",
    "Last Name",
    "Position / Job Title",
    "Company",
    "Location",
    "Phone Number",
    "Email Address",
]

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

COLUMN_WIDTHS = [15, 15, 25, 25, 20, 20, 30]


class ExportService:
    """Generates Excel exports from lead data."""

    def generate_excel(self, leads: List[LeadResponse]) -> bytes:
        try:
            with log_duration(logger, "Excel export", lead_count=len(leads)):
                wb = Workbook()
                ws = wb.active
                ws.title = "Leads"

                for col_idx, header in enumerate(HEADERS, 1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = HEADER_ALIGNMENT

                ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
                ws.freeze_panes = "A2"

                for col_idx, width in enumerate(COLUMN_WIDTHS, 1):
                    ws.column_dimensions[get_column_letter(col_idx)].width = width

                for row_idx, lead in enumerate(leads, 2):
                    values = [
                        lead.first_name,
                        lead.last_name,
                        lead.position,
                        lead.company,
                        lead.location,
                        lead.phone,
                        lead.email,
                    ]
                    for col_idx, value in enumerate(values, 1):
                        cell = ws.cell(row=row_idx, column=col_idx, value=value or "")
                        cell.alignment = CELL_ALIGNMENT

                buffer = BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                return buffer.read()

        except Exception as e:
            logger.error("Excel export failed: %s", e)
            raise ExportError(f"Failed to generate Excel: {e}") from e
